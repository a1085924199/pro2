#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ui_window.py - 主窗口 v3.0
重构要点：
  - FastGPT 配置移入独立对话框（高级设置）
  - 专用单据：左图右字段并排，字段卡片视图
  - 通用表格：HTML渲染预览 + 结构 + 汇总
  - 进度条改为醒目横幅
"""
import os, json
from datetime import datetime

from PyQt5.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QStackedWidget,
    QPushButton, QLabel, QLineEdit, QCheckBox, QProgressBar, QDialog,
    QTabWidget, QTableWidget, QTableWidgetItem, QTextEdit, QTextBrowser,
    QListWidget, QListWidgetItem, QGroupBox, QGridLayout, QFormLayout,
    QFileDialog, QMessageBox, QSplitter, QHeaderView, QFrame,
    QScrollArea, QSizePolicy, QDialogButtonBox
)
from PyQt5.QtCore import Qt, QTimer, QPropertyAnimation, QEasingCurve
from PyQt5.QtGui import QPixmap, QFont, QColor, QColor
from openpyxl import Workbook
from openpyxl.styles import Font as XLFont, Alignment, PatternFill, Border, Side
from ocr_core import BatchOCRWorker, TableStructureWorker

# ============================================================================
# 样式
# ============================================================================
SS = """
QMainWindow,QWidget{background:#f5f6fa;font-family:'Microsoft YaHei',sans-serif;font-size:13px;}
QPushButton{background:#1677ff;color:white;border:none;border-radius:6px;
    padding:6px 16px;font-size:13px;font-weight:600;min-height:32px;}
QPushButton:hover{background:#4096ff;}
QPushButton:pressed{background:#0958d9;}
QPushButton:disabled{background:#e4e6ea;color:#adb5bd;}
QPushButton#sec{background:white;color:#495057;border:1px solid #dee2e6;}
QPushButton#sec:hover{border-color:#4096ff;color:#1677ff;background:#f0f7ff;}
QPushButton#sec:disabled{background:#f8f9fa;color:#adb5bd;}
QPushButton#warn{background:#fa5252;}
QPushButton#warn:hover{background:#ff6b6b;}
QPushButton#inactive{background:#f8f9fa;color:#adb5bd;border:1px dashed #dee2e6;font-weight:400;}
QPushButton#active_form{background:#e7f5ff;color:#1677ff;
    border:2px solid #1677ff;border-radius:8px;font-size:13px;font-weight:700;}
QGroupBox{font-weight:700;border:1px solid #e9ecef;border-radius:10px;
    margin-top:16px;padding-top:16px;background:white;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04);}
QGroupBox::title{subcontrol-origin:margin;left:14px;padding:0 8px;
    color:#343a40;background:white;}
QTabWidget#mt::pane{border:none;background:#f5f6fa;}
QTabWidget#mt>QTabBar{background:transparent;}
QTabWidget#mt>QTabBar::tab{background:white;color:#6c757d;
    padding:12px 32px;font-size:14px;font-weight:700;
    border:1px solid #e9ecef;border-bottom:none;
    border-radius:8px 8px 0 0;margin-right:4px;}
QTabWidget#mt>QTabBar::tab:selected{background:#1677ff;color:white;border-color:#1677ff;}
QTabWidget#mt>QTabBar::tab:hover:!selected{background:#e7f5ff;color:#1677ff;}
QTabWidget#sub::pane{border:1px solid #e9ecef;border-radius:8px;background:white;}
QTabWidget#sub>QTabBar::tab{background:#f8f9fa;color:#6c757d;
    padding:7px 18px;border-radius:6px 6px 0 0;margin-right:2px;font-weight:600;}
QTabWidget#sub>QTabBar::tab:selected{background:white;color:#1677ff;
    border-bottom:2px solid #1677ff;}
QTableWidget{gridline-color:#f1f3f5;background:white;
    border:1px solid #e9ecef;border-radius:8px;
    selection-background-color:#e7f5ff;selection-color:#212529;}
QTableWidget::item{padding:6px 10px;border-bottom:1px solid #f8f9fa;}
QTableWidget::item:alternate{background:#fdfdfe;}
QHeaderView::section{background:#f8f9fa;padding:8px 12px;
    border:none;border-bottom:2px solid #e9ecef;
    font-weight:700;color:#495057;font-size:12px;}
QListWidget{border:1px solid #e9ecef;border-radius:8px;padding:4px;background:white;}
QListWidget::item{padding:8px 12px;border-radius:6px;margin-bottom:2px;color:#343a40;}
QListWidget::item:selected{background:#e7f5ff;color:#1677ff;font-weight:600;}
QListWidget::item:hover{background:#f8f9fa;}
QLineEdit{padding:7px 12px;border:1px solid #dee2e6;
    border-radius:6px;background:white;color:#212529;}
QLineEdit:focus{border:1.5px solid #1677ff;background:#fafcff;}
QProgressBar{border:none;border-radius:3px;background:#e9ecef;height:6px;text-align:center;}
QProgressBar::chunk{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #1677ff,stop:1 #69b1ff);border-radius:3px;}
QTextEdit,QTextBrowser{background:#f8f9fa;border:1px solid #e9ecef;border-radius:6px;color:#212529;}
QScrollBar:vertical{width:6px;background:transparent;border-radius:3px;}
QScrollBar::handle:vertical{background:#ced4da;border-radius:3px;min-height:20px;}
QScrollBar::add-line:vertical,QScrollBar::sub-line:vertical{height:0;}
"""

# ============================================================================
# 高级设置对话框（FastGPT配置）
# ============================================================================
class AdvancedSettingsDialog(QDialog):
    def __init__(self, config, parent=None):
        super().__init__(parent)
        self.setWindowTitle("高级设置 - FastGPT 配置")
        self.setMinimumWidth(500)
        self.setStyleSheet(SS)
        layout = QVBoxLayout(self)
        layout.setSpacing(16); layout.setContentsMargins(24,24,24,24)

        desc = QLabel("配置 FastGPT 智能解析接口（可选）。未配置时使用本地规则提取字段。")
        desc.setWordWrap(True)
        desc.setStyleSheet("color:#6c757d;background:#f8f9fa;padding:10px;border-radius:6px;")
        layout.addWidget(desc)

        form = QFormLayout(); form.setSpacing(12)
        self.edit_url   = QLineEdit(config.get("api_url",""))
        self.edit_key   = QLineEdit(config.get("api_key",""))
        self.edit_key.setEchoMode(QLineEdit.Password)
        self.edit_appid = QLineEdit(config.get("appid",""))
        self.cb_enable  = QCheckBox("启用 FastGPT 智能解析")
        self.cb_enable.setChecked(config.get("enabled", False))
        form.addRow("API 地址:", self.edit_url)
        form.addRow("API 密钥:", self.edit_key)
        form.addRow("应用 ID:",  self.edit_appid)
        form.addRow("",          self.cb_enable)
        layout.addLayout(form)

        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

    def get_config(self):
        return {
            "api_url": self.edit_url.text(),
            "api_key": self.edit_key.text(),
            "appid":   self.edit_appid.text(),
            "enabled": self.cb_enable.isChecked()
        }


# ============================================================================
# 进度横幅
# ============================================================================
class ProgressBanner(QFrame):
    """识别中醒目进度横幅"""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFixedHeight(44)
        self.setStyleSheet("background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #1677ff,stop:1 #69b1ff);border-radius:8px;")
        hl = QHBoxLayout(self); hl.setContentsMargins(16,0,16,0); hl.setSpacing(12)
        self.lbl_msg = QLabel("正在处理..."); self.lbl_msg.setStyleSheet("color:white;font-weight:600;")
        self.bar = QProgressBar()
        self.bar.setStyleSheet("QProgressBar{background:rgba(255,255,255,0.3);border-radius:3px;height:8px;}QProgressBar::chunk{background:white;border-radius:3px;}")
        self.bar.setFixedHeight(8); self.bar.setTextVisible(False)
        self.lbl_pct = QLabel("0%"); self.lbl_pct.setStyleSheet("color:white;font-weight:700;min-width:36px;")
        hl.addWidget(self.lbl_msg, stretch=1)
        hl.addWidget(self.bar, stretch=2)
        hl.addWidget(self.lbl_pct)
        self.hide()

    def update(self, val, msg):
        self.bar.setValue(val); self.lbl_msg.setText(msg); self.lbl_pct.setText(f"{val}%")
        self.show()

    def done(self):
        self.hide(); self.bar.setValue(0)


# ============================================================================
# 字段卡片组件
# ============================================================================
class FieldCard(QFrame):
    """单个字段的卡片展示"""
    def __init__(self, label, value="", parent=None):
        super().__init__(parent)
        self.setStyleSheet(
            "QFrame{background:white;border:1px solid #e9ecef;border-radius:8px;padding:2px;}"
            "QFrame:hover{border-color:#74c0fc;}"
        )
        vl = QVBoxLayout(self); vl.setContentsMargins(12,10,12,10); vl.setSpacing(4)
        lbl = QLabel(label); lbl.setStyleSheet("color:#6c757d;font-size:11px;font-weight:600;")
        self.val_label = QLabel(value if value else "—")
        self.val_label.setWordWrap(True)
        self.val_label.setStyleSheet("color:#212529;font-size:14px;font-weight:600;" if value
                                      else "color:#adb5bd;font-size:14px;")
        vl.addWidget(lbl); vl.addWidget(self.val_label)

    def set_value(self, value):
        self.val_label.setText(value if value else "—")
        self.val_label.setStyleSheet("color:#212529;font-size:14px;font-weight:600;" if value
                                      else "color:#adb5bd;font-size:14px;")


# ============================================================================
# 工具函数
# ============================================================================
def _btn(text, slot, enabled=True, oid=None, h=34):
    b = QPushButton(text); b.setMinimumHeight(h)
    b.clicked.connect(slot); b.setEnabled(enabled)
    if oid: b.setObjectName(oid)
    return b

def _ro_edit(fs=10):
    t = QTextEdit(); t.setReadOnly(True); t.setFont(QFont("Consolas", fs)); return t

def _add_tab(tabs, w, title):
    c = QWidget(); vl = QVBoxLayout(c)
    vl.setContentsMargins(6,6,6,6); vl.addWidget(w); tabs.addTab(c, title)
# Part B: MainWindow skeleton + Tab1 layout
class MainWindow(QMainWindow):
    APP_VERSION = "3.0.0"
    def __init__(self):
        super().__init__()
        self.setWindowTitle("智·链 OCR 平台")
        self.setGeometry(80, 80, 1440, 900)
        self.setStyleSheet(SS)
        self._r_fgpt = {"api_url":"","api_key":"","appid":"","enabled":False}
        self._r_images=[]; self._r_results=[]; self._r_idx=-1; self._r_worker=None
        root = QWidget(); self.setCentralWidget(root)
        rl = QVBoxLayout(root); rl.setContentsMargins(0,0,0,0); rl.setSpacing(0)
        rl.addWidget(self._hdr())
        rl.addWidget(self._main_tabs(), stretch=1)
        rl.addWidget(self._footer())
        self._load_fgpt_config()

    def _hdr(self):
        f = QFrame(); f.setFixedHeight(52)
        f.setStyleSheet("background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #1677ff,stop:1 #0ea5e9);")
        hl = QHBoxLayout(f); hl.setContentsMargins(20,0,20,0)
        t = QLabel("单据OCR识别系统"); t.setStyleSheet("color:white;font-size:17px;font-weight:700;")
        hl.addWidget(t); hl.addStretch()
        v = QLabel(f"v{self.APP_VERSION}  ·  PaddleOCR 3.x"); v.setStyleSheet("color:rgba(255,255,255,0.7);font-size:12px;")
        hl.addWidget(v); return f

    def _main_tabs(self):
        self.mtabs = QTabWidget(); self.mtabs.setObjectName("mt")
        self.mtabs.addTab(self._build_forms_tab(), "📋  专用单据识别")
        self.mtabs.addTab(self._build_table_tab(), "📊  通用表格识别")
        return self.mtabs

    def _footer(self):
        f = QFrame(); f.setFixedHeight(28); f.setStyleSheet("background:white;border-top:1px solid #e9ecef;")
        fl = QHBoxLayout(f); fl.setContentsMargins(16,0,16,0)
        self.lbl_status = QLabel("就绪"); self.lbl_status.setStyleSheet("color:#6c757d;font-size:12px;")
        fl.addWidget(self.lbl_status); fl.addStretch(); return f

    def _build_forms_tab(self):
        page = QWidget()
        pl = QVBoxLayout(page); pl.setContentsMargins(16,12,16,12); pl.setSpacing(10)
        # 类型选择栏
        tbar = QHBoxLayout(); tbar.setSpacing(10)
        for label, active in [("🔧 维修器材调修单",True),("↩ 返修件返修卡",False),("✈ 航材入库单",False),("📦 航材出库单",False)]:
            if active:
                b = QPushButton(label); b.setMinimumHeight(44); b.setObjectName("active_form"); b.clicked.connect(lambda:None)
            else:
                b = QPushButton(label); b.setMinimumHeight(44); b.setObjectName("inactive")
                _l=label; b.clicked.connect(lambda _,l=_l: QMessageBox.information(self,"提示",f"【{l}】功能开发中。"))
            tbar.addWidget(b)
        pl.addLayout(tbar)
        # 工具栏
        tb = QHBoxLayout(); tb.setSpacing(8)
        self.r_btn_sel  = _btn("📁 选择图片",  self._r_select)
        self.r_btn_clr  = _btn("🗑 清空",        self._r_clear,  enabled=False, oid="sec")
        self.r_btn_run  = _btn("▶ 开始识别",    self._r_start,  enabled=False)
        self.r_btn_stop = _btn("⏹ 停止",         self._r_stop,   enabled=False, oid="warn")
        self.r_btn_adv  = _btn("⚙ 高级设置",    self._r_advanced, oid="sec")
        self.r_btn_xl   = _btn("📊 导出Excel",  self._r_excel,  enabled=False, oid="sec")
        self.r_btn_js   = _btn("💾 保存JSON",   self._r_json,   enabled=False, oid="sec")
        for b in [self.r_btn_sel,self.r_btn_clr]: tb.addWidget(b)
        tb.addSpacing(12)
        for b in [self.r_btn_run,self.r_btn_stop]: tb.addWidget(b)
        tb.addSpacing(12); tb.addWidget(self.r_btn_adv); tb.addSpacing(12)
        for b in [self.r_btn_xl,self.r_btn_js]: tb.addWidget(b)
        tb.addStretch()
        self.r_fgpt_badge = QLabel("⚡ FastGPT已启用")
        self.r_fgpt_badge.setStyleSheet("background:#ebfbee;color:#2f9e44;border:1px solid #b2f2bb;border-radius:10px;padding:3px 10px;font-size:11px;font-weight:600;")
        self.r_fgpt_badge.hide(); tb.addWidget(self.r_fgpt_badge)
        pl.addLayout(tb)
        self.r_banner = ProgressBanner(); pl.addWidget(self.r_banner)
        sp = QSplitter(Qt.Horizontal); sp.setHandleWidth(8)
        sp.addWidget(self._build_repair_left()); sp.addWidget(self._build_repair_right())
        sp.setStretchFactor(0,4); sp.setStretchFactor(1,6)
        pl.addWidget(sp, stretch=1)
        return page

    def _build_repair_left(self):
        panel = QWidget()
        pl = QVBoxLayout(panel); pl.setContentsMargins(0,0,6,0); pl.setSpacing(8)
        lg = QGroupBox("图片列表"); ll = QVBoxLayout(lg)
        h = QHBoxLayout()
        self.r_lbl_cnt = QLabel("共 0 张"); self.r_lbl_cnt.setStyleSheet("color:#1677ff;font-weight:700;")
        h.addWidget(self.r_lbl_cnt); h.addStretch(); ll.addLayout(h)
        self.r_img_list = QListWidget(); self.r_img_list.setMinimumHeight(140)
        self.r_img_list.itemClicked.connect(self._r_on_click); ll.addWidget(self.r_img_list)
        pl.addWidget(lg)
        pg = QGroupBox("图片预览"); pgl = QVBoxLayout(pg)
        self.r_lbl_prev = QLabel("选择图片后预览")
        self.r_lbl_prev.setAlignment(Qt.AlignCenter); self.r_lbl_prev.setMinimumHeight(220)
        self.r_lbl_prev.setStyleSheet("border:2px dashed #dee2e6;border-radius:10px;background:#f8f9fa;color:#adb5bd;")
        pgl.addWidget(self.r_lbl_prev)
        self.r_lbl_info = QLabel(""); self.r_lbl_info.setAlignment(Qt.AlignCenter)
        self.r_lbl_info.setStyleSheet("color:#adb5bd;font-size:11px;margin-top:4px;")
        pgl.addWidget(self.r_lbl_info)
        pl.addWidget(pg, stretch=1); return panel

    def _build_repair_right(self):
        panel = QWidget(); pl = QVBoxLayout(panel); pl.setContentsMargins(6,0,0,0)
        tabs = QTabWidget(); tabs.setObjectName("sub")
        # 字段卡片
        sa = QScrollArea(); sa.setWidgetResizable(True)
        sa.setStyleSheet("QScrollArea{border:none;background:transparent;}")
        fw = QWidget(); self.r_form_gl = QGridLayout(fw)
        self.r_form_gl.setSpacing(10); self.r_form_gl.setContentsMargins(8,8,8,8)
        self._field_cards = {}
        for fname,row,col,span in [("调修单号",0,0,1),("装备型号",0,1,1),("器材名称",1,0,1),("型（图）号",1,1,1),("器件编号",2,0,1),("进厂时间",2,1,1),("邮寄地址",3,0,2)]:
            card = FieldCard(fname); self.r_form_gl.addWidget(card,row,col,1,span); self._field_cards[fname]=card
        self.r_ocr_summary = QLabel("")
        self.r_ocr_summary.setWordWrap(True)
        self.r_ocr_summary.setStyleSheet("color:#6c757d;font-size:11px;background:#f8f9fa;border:1px solid #e9ecef;border-radius:6px;padding:6px 12px;")
        self.r_form_gl.addWidget(self.r_ocr_summary,4,0,1,2)
        sa.setWidget(fw); _add_tab(tabs, sa, "📋 识别结果")
        # 批量汇总
        cols=["图片名称","调修单号","装备型号","器材名称","型（图）号","器件编号","邮寄地址","进厂时间"]
        self.r_tbl_all = QTableWidget(0,len(cols)); self.r_tbl_all.setHorizontalHeaderLabels(cols)
        self.r_tbl_all.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.r_tbl_all.setAlternatingRowColors(True)
        _add_tab(tabs, self.r_tbl_all, "📊 批量汇总")
        self.r_txt_json = _ro_edit(10); _add_tab(tabs, self.r_txt_json, "{ } JSON")
        pl.addWidget(tabs, stretch=1); return panel
    # --- FastGPT 配置 ---
    def _r_advanced(self):
        dlg = AdvancedSettingsDialog(self._r_fgpt, self)
        if dlg.exec_() == QDialog.Accepted:
            self._r_fgpt = dlg.get_config()
            self._save_fgpt_config()
            self.r_fgpt_badge.setVisible(self._r_fgpt.get("enabled",False))

    def _load_fgpt_config(self):
        p = os.path.join(os.path.dirname(os.path.abspath(__file__)),"config.json")
        if os.path.exists(p):
            try:
                c = json.load(open(p,encoding="utf-8"))
                self._r_fgpt.update(c)
                self.r_fgpt_badge.setVisible(self._r_fgpt.get("enabled",False))
            except Exception as e: print(f"加载配置失败:{e}")

    def _save_fgpt_config(self):
        p = os.path.join(os.path.dirname(os.path.abspath(__file__)),"config.json")
        try: json.dump(self._r_fgpt,open(p,"w",encoding="utf-8"),ensure_ascii=False,indent=2)
        except Exception as e: print(f"保存配置失败:{e}")

    # --- 图片操作 ---
    def _r_select(self):
        paths,_ = QFileDialog.getOpenFileNames(self,"选择图片","","图片 (*.jpg *.jpeg *.png *.bmp *.tif *.tiff)")
        for p in paths:
            if p not in self._r_images:
                self._r_images.append(p)
                it = QListWidgetItem(os.path.basename(p)); it.setData(Qt.UserRole,p)
                self.r_img_list.addItem(it)
        if paths:
            self._r_update_cnt(); self.r_btn_run.setEnabled(True); self.r_btn_clr.setEnabled(True)
            self.r_img_list.setCurrentRow(0); self._r_on_click(self.r_img_list.item(0))

    def _r_clear(self):
        self._r_images=[]; self._r_results=[]; self._r_idx=-1
        self.r_img_list.clear(); self.r_tbl_all.setRowCount(0); self.r_txt_json.clear()
        for c in self._field_cards.values(): c.set_value("")
        self.r_ocr_summary.setText("")
        self.r_lbl_prev.clear(); self.r_lbl_prev.setText("选择图片后预览"); self.r_lbl_info.setText("")
        self._r_update_cnt()
        for b in (self.r_btn_run,self.r_btn_clr,self.r_btn_xl,self.r_btn_js): b.setEnabled(False)
        self.lbl_status.setText("就绪")

    def _r_update_cnt(self):
        n=len(self._r_images); self.r_lbl_cnt.setText(f"共 {n} 张")
        self.lbl_status.setText(f"已选择 {n} 张" if n else "就绪")

    def _r_on_click(self, item):
        if not item: return
        path=item.data(Qt.UserRole)
        self._r_idx=self._r_images.index(path) if path in self._r_images else -1
        px=QPixmap(path)
        if not px.isNull():
            self.r_lbl_prev.setPixmap(px.scaled(420,280,Qt.KeepAspectRatio,Qt.SmoothTransformation))
            self.r_lbl_info.setText(f"{px.width()}x{px.height()}  {os.path.getsize(path)/1024:.1f}KB")
        else:
            self.r_lbl_prev.setText("无法加载图片"); self.r_lbl_info.setText("")
        idx=self._r_idx
        if 0<=idx<len(self._r_results) and self._r_results[idx]: self._r_show(self._r_results[idx])

    # --- OCR 流程 ---
    def _r_start(self):
        if not self._r_images: return
        self._save_fgpt_config()
        for b in (self.r_btn_run,self.r_btn_sel,self.r_btn_clr,self.r_btn_xl,self.r_btn_js): b.setEnabled(False)
        self.r_btn_stop.setEnabled(True)
        self.r_banner.update(0,"正在初始化 OCR 引擎...")
        self.lbl_status.setText("正在处理...")
        self._r_results=[None]*len(self._r_images)
        self._r_worker=BatchOCRWorker(self._r_images,self._r_fgpt.get("enabled",False),self._r_fgpt)
        self._r_worker.progress.connect(self._r_on_prog)
        self._r_worker.finished.connect(self._r_on_done)
        self._r_worker.error.connect(self._r_on_err)
        self._r_worker.image_finished.connect(self._r_on_img_done)
        self._r_worker.start()

    def _r_stop(self):
        if self._r_worker: self._r_worker.stop()
        self.r_btn_stop.setEnabled(False); self.lbl_status.setText("正在停止...")

    def _r_on_prog(self,val,msg,cur,total):
        overall=int((cur-1)*100/total+val/total)
        self.r_banner.update(overall,f"{msg}  ({cur}/{total})")
        self.lbl_status.setText(msg)

    def _r_on_img_done(self,result,idx):
        self._r_results[idx]=result
        item=self.r_img_list.item(idx)
        if result.get("success"): item.setForeground(Qt.darkGreen); item.setText(f"\u2713 {result['image_name']}")
        else: item.setForeground(Qt.red); item.setText(f"\u2717 {result['image_name']}")
        if idx==self._r_idx: self._r_show(result)
        self._r_refresh_all()

    def _r_on_done(self,results):
        for b in (self.r_btn_run,self.r_btn_sel,self.r_btn_clr,self.r_btn_xl,self.r_btn_js): b.setEnabled(True)
        self.r_btn_stop.setEnabled(False); self.r_banner.done()
        self.lbl_status.setText("处理完成")
        ok=sum(1 for r in results if r.get("success"))
        QMessageBox.information(self,"完成",f"处理完成！\n成功: {ok}/{len(results)} 张")

    def _r_on_err(self,msg):
        for b in (self.r_btn_run,self.r_btn_sel,self.r_btn_clr): b.setEnabled(True)
        self.r_btn_stop.setEnabled(False); self.r_banner.done()
        self.lbl_status.setText("处理失败"); QMessageBox.critical(self,"错误",msg)

    def _r_show(self,result):
        if result.get("success"):
            fields=result.get("extracted_fields",{})
            for k,card in self._field_cards.items(): card.set_value(fields.get(k,""))
            ocr_list=result.get("ocr_results",[])
            self.r_ocr_summary.setText(f"OCR 识别 {len(ocr_list)} 条文本  ·  图片: {result.get('image_name','')}  ·  {result.get('timestamp','')}")
            self.r_txt_json.setText(json.dumps({"识别状态":"成功","图片名称":result.get("image_name",""),
                "识别时间":result.get("timestamp",""),"提取字段":fields,"OCR结果":ocr_list},ensure_ascii=False,indent=2))
        else:
            for card in self._field_cards.values(): card.set_value("")
            self.r_ocr_summary.setText(f"\u274c 处理失败: {result.get('error','未知')}")
            self.r_txt_json.setText(json.dumps(result,ensure_ascii=False,indent=2))

    def _r_refresh_all(self):
        valid=[r for r in self._r_results if r and r.get("success")]
        self.r_tbl_all.setRowCount(len(valid))
        keys=["调修单号","装备型号","器材名称","型（图）号","器件编号","邮寄地址","进厂时间"]
        for row,r in enumerate(valid):
            f=r.get("extracted_fields",{})
            self.r_tbl_all.setItem(row,0,QTableWidgetItem(r.get("image_name","")))
            for col,k in enumerate(keys,1): self.r_tbl_all.setItem(row,col,QTableWidgetItem(f.get(k,"")))

    def _r_excel(self):
        if not any(r and r.get("success") for r in self._r_results):
            QMessageBox.warning(self,"警告","没有可导出的结果"); return
        path,_=QFileDialog.getSaveFileName(self,"保存Excel",f"调修单识别_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx","Excel (*.xlsx)")
        if not path: return
        try:
            wb=Workbook(); ws=wb.active; ws.title="识别结果"
            hf=PatternFill(start_color="1677FF",end_color="1677FF",fill_type="solid")
            hfont=XLFont(bold=True,color="FFFFFF",size=11)
            bd=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
            hdrs=["图片名称","调修单号","装备型号","器材名称","型（图）号","器件编号","邮寄地址","进厂时间","识别时间"]
            for c,h in enumerate(hdrs,1):
                cell=ws.cell(1,c,h); cell.fill=hf; cell.font=hfont
                cell.alignment=Alignment(horizontal='center',vertical='center'); cell.border=bd
            for row,r in enumerate([x for x in self._r_results if x and x.get("success")],2):
                f=r.get("extracted_fields",{})
                vals=[r.get("image_name",""),f.get("调修单号",""),f.get("装备型号",""),f.get("器材名称",""),f.get("型（图）号",""),f.get("器件编号",""),f.get("邮寄地址",""),f.get("进厂时间",""),r.get("timestamp","")]
                for c,v in enumerate(vals,1):
                    cell=ws.cell(row,c,v); cell.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True); cell.border=bd
            for c in range(1,len(hdrs)+1): ws.column_dimensions[chr(64+c)].width=20
            wb.save(path); QMessageBox.information(self,"成功",f"已保存:\n{path}")
        except Exception as e: QMessageBox.critical(self,"错误",f"导出失败:{e}")

    def _r_json(self):
        if not self._r_results: QMessageBox.warning(self,"警告","没有结果"); return
        d=QFileDialog.getExistingDirectory(self,"选择保存目录",".")
        if not d: return
        ts=datetime.now().strftime("%Y%m%d_%H%M%S")
        out={"导出时间":datetime.now().strftime("%Y-%m-%d %H:%M:%S"),"总图片数":len(self._r_images),
             "成功数":sum(1 for r in self._r_results if r and r.get("success")),"结果列表":[r for r in self._r_results if r]}
        p=os.path.join(d,f"调修单识别_{ts}.json")
        try:
            json.dump(out,open(p,"w",encoding="utf-8"),ensure_ascii=False,indent=2)
            QMessageBox.information(self,"成功",f"已保存:\n{p}")
        except Exception as e: QMessageBox.critical(self,"错误",f"保存失败:{e}")
    def _t_json(self):
        if not self._t_results: QMessageBox.warning(self,"警告","没有结果"); return
        d=QFileDialog.getExistingDirectory(self,"选择保存目录",".")
        if not d: return
        ts=datetime.now().strftime("%Y%m%d_%H%M%S")
        p=os.path.join(d,f"表格识别_{ts}.json")
        try:
            json.dump(self._t_results,open(p,"w",encoding="utf-8"),ensure_ascii=False,indent=2)
            QMessageBox.information(self,"成功",f"已保存:\n{p}")
        except Exception as e: QMessageBox.critical(self,"错误",f"保存失败:{e}")

    def _t_excel(self):
        if not any(r.get("success") for r in self._t_results):
            QMessageBox.warning(self,"警告","没有可导出的结果"); return
        path,_=QFileDialog.getSaveFileName(self,"保存Excel",
            f"表格识别_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx","Excel (*.xlsx)")
        if not path: return
        try:
            wb=Workbook(); wb.remove(wb.active)
            hf=PatternFill(start_color="1677FF",end_color="1677FF",fill_type="solid")
            hfont=XLFont(bold=True,color="FFFFFF",size=11)
            bd=Border(left=Side(style='thin'),right=Side(style='thin'),
                      top=Side(style='thin'),bottom=Side(style='thin'))
            for r in self._t_results:
                if not r.get("success"): continue
                for tbl in r.get("tables",[]):
                    sname=f"{r['file_name'][:14]}_T{tbl['table_idx']+1}"
                    ws=wb.create_sheet(title=sname[:31])
                    cells=tbl.get("cells",[])
                    if not cells: ws.cell(1,1,"无单元格数据"); continue
                    for c in cells:
                        rs=c.get("row_span",1); cs2=c.get("col_span",1)
                        er=c["row"]+1; ec=c["col"]+1
                        cell=ws.cell(er,ec,c.get("text",""))
                        cell.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
                        cell.border=bd
                        if rs>1 or cs2>1:
                            try: ws.merge_cells(start_row=er,start_column=ec,end_row=er+rs-1,end_column=ec+cs2-1)
                            except: pass
                    mc=max(c["col"]+c.get("col_span",1) for c in cells)
                    for ci in range(1,mc+1):
                        ws.column_dimensions[chr(64+min(ci,26))].width=18
            wb.save(path); QMessageBox.information(self,"成功",f"已保存:\n{path}")
        except Exception as e: QMessageBox.critical(self,"错误",f"导出失败:{e}")

    def _build_table_tab(self):
        self._t_files=[]; self._t_results=[]; self._t_worker=None
        page = QWidget()
        pl = QVBoxLayout(page); pl.setContentsMargins(16,12,16,12); pl.setSpacing(10)
        info = QLabel("使用 PP-StructureV3 识别任意格式的表格文档（图片或PDF），还原表格逻辑结构。")
        info.setWordWrap(True)
        info.setStyleSheet("background:#e7f5ff;color:#1677ff;padding:10px 14px;border-radius:8px;font-size:12px;border:1px solid #74c0fc;")
        pl.addWidget(info)
        tb = QHBoxLayout(); tb.setSpacing(8)
        self.t_btn_sel  = _btn("选择文件", self._t_select)
        self.t_btn_clr  = _btn("清空",   self._t_clear,  enabled=False, oid="sec")
        self.t_btn_run  = _btn("开始识别", self._t_start,  enabled=False)
        self.t_btn_stop = _btn("停止",    self._t_stop,   enabled=False, oid="warn")
        self.t_btn_xl   = _btn("导出Excel",self._t_excel,  enabled=False, oid="sec")
        self.t_btn_js   = _btn("导出JSON", self._t_json,   enabled=False, oid="sec")
        for b in [self.t_btn_sel,self.t_btn_clr]: tb.addWidget(b)
        tb.addSpacing(12)
        for b in [self.t_btn_run,self.t_btn_stop]: tb.addWidget(b)
        tb.addSpacing(12)
        for b in [self.t_btn_xl,self.t_btn_js]: tb.addWidget(b)
        tb.addStretch(); pl.addLayout(tb)
        self.t_banner = ProgressBanner(); pl.addWidget(self.t_banner)
        sp = QSplitter(Qt.Horizontal); sp.setHandleWidth(8)
        sp.addWidget(self._build_table_left()); sp.addWidget(self._build_table_right())
        sp.setStretchFactor(0,3); sp.setStretchFactor(1,7)
        pl.addWidget(sp, stretch=1)
        return page

    def _build_table_left(self):
        panel = QWidget()
        pl = QVBoxLayout(panel); pl.setContentsMargins(0,0,6,0); pl.setSpacing(8)
        lg = QGroupBox("文件列表"); ll = QVBoxLayout(lg)
        h = QHBoxLayout()
        self.t_lbl_cnt = QLabel("共 0 个"); self.t_lbl_cnt.setStyleSheet("color:#1677ff;font-weight:700;")
        h.addWidget(self.t_lbl_cnt); h.addStretch(); ll.addLayout(h)
        self.t_file_list = QListWidget(); self.t_file_list.setMinimumHeight(160)
        self.t_file_list.itemClicked.connect(self._t_on_file_click); ll.addWidget(self.t_file_list)
        pl.addWidget(lg)
        tg = QGroupBox("识别到的表格"); tgl = QVBoxLayout(tg)
        self.t_tbl_list = QListWidget(); self.t_tbl_list.setMinimumHeight(100)
        self.t_tbl_list.itemClicked.connect(self._t_on_tbl_click); tgl.addWidget(self.t_tbl_list)
        pl.addWidget(tg)
        sg = QGroupBox("汇总统计"); sgl = QVBoxLayout(sg)
        self.t_lbl_summary = QLabel("暂无数据")
        self.t_lbl_summary.setWordWrap(True)
        self.t_lbl_summary.setStyleSheet("color:#495057;font-size:12px;")
        sgl.addWidget(self.t_lbl_summary); pl.addWidget(sg)
        return panel

    def _build_table_right(self):
        panel = QWidget(); pl = QVBoxLayout(panel); pl.setContentsMargins(6,0,0,0)
        tabs = QTabWidget(); tabs.setObjectName("sub")
        self.t_browser = QTextBrowser()
        self.t_browser.setFont(QFont("Microsoft YaHei", 11))
        self.t_browser.setStyleSheet("background:white;border:1px solid #e9ecef;border-radius:8px;padding:8px;")
        _add_tab(tabs, self.t_browser, "表格预览")
        self.t_tbl_cells = QTableWidget(0, 7)
        self.t_tbl_cells.setHorizontalHeaderLabels(["行","列","行跨","列跨","内容","置信度","坐标(x1,y1,x2,y2)"])
        self.t_tbl_cells.horizontalHeader().setSectionResizeMode(4, QHeaderView.Stretch)
        self.t_tbl_cells.horizontalHeader().setSectionResizeMode(6, QHeaderView.ResizeToContents)
        self.t_tbl_cells.setAlternatingRowColors(True)
        _add_tab(tabs, self.t_tbl_cells, "单元格结构")
        self.t_txt_json = _ro_edit(9); _add_tab(tabs, self.t_txt_json, "JSON")
        pl.addWidget(tabs, stretch=1); return panel

    def _t_select(self):
        paths,_ = QFileDialog.getOpenFileNames(self,"选择文件","","支持的文件 (*.jpg *.jpeg *.png *.bmp *.tif *.tiff *.pdf)")
        for p in paths:
            if p not in self._t_files:
                self._t_files.append(p)
                it = QListWidgetItem(os.path.basename(p)); it.setData(Qt.UserRole,p)
                self.t_file_list.addItem(it)
        if paths: self._t_update_cnt(); self.t_btn_run.setEnabled(True); self.t_btn_clr.setEnabled(True)

    def _t_clear(self):
        self._t_files=[]; self._t_results=[]
        self.t_file_list.clear(); self.t_tbl_list.clear()
        self.t_tbl_cells.setRowCount(0); self.t_txt_json.clear()
        self.t_browser.clear(); self.t_lbl_summary.setText("暂无数据")
        self._t_update_cnt()
        for b in (self.t_btn_run,self.t_btn_clr,self.t_btn_xl,self.t_btn_js): b.setEnabled(False)
        self.lbl_status.setText("就绪")

    def _t_update_cnt(self):
        n=len(self._t_files); self.t_lbl_cnt.setText(f"共 {n} 个")
        self.lbl_status.setText(f"已选择 {n} 个文件" if n else "就绪")

    def _t_on_file_click(self, item):
        if not item: return
        path=item.data(Qt.UserRole)
        result=next((r for r in self._t_results if r.get("file_path")==path),None)
        self.t_tbl_list.clear()
        if not result: return
        if result.get("success"):
            for tbl in result.get("tables",[]):
                cells=tbl.get("cells",[])
                nr=max((c["row"]+c.get("row_span",1) for c in cells),default=0)
                nc=max((c["col"]+c.get("col_span",1) for c in cells),default=0)
                it=QListWidgetItem(f"表格 {tbl['table_idx']+1}  ({nr}行 x {nc}列)")
                it.setData(Qt.UserRole,tbl); self.t_tbl_list.addItem(it)
            if self.t_tbl_list.count()>0:
                self.t_tbl_list.setCurrentRow(0); self._t_on_tbl_click(self.t_tbl_list.item(0))
        else:
            self.t_browser.setHtml(f"<p style='color:red'>识别失败: {result.get('error','')}</p>")

    def _t_on_tbl_click(self, item):
        if not item: return
        tbl=item.data(Qt.UserRole)
        if not tbl: return
        html=tbl.get("html","")
        if html:
            styled=("<style>table{border-collapse:collapse;width:100%;font-family:'Microsoft YaHei',sans-serif;font-size:13px;}"
                    "td,th{border:1px solid #dee2e6;padding:8px 12px;}th{background:#f8f9fa;font-weight:700;}"
                    "tr:nth-child(even){background:#f8f9fa;}</style>") + html
            self.t_browser.setHtml(styled)
        else:
            self.t_browser.setHtml("<p style='color:#6c757d;padding:20px'>该表格无HTML输出。</p>")
        cells=tbl.get("cells",[])
        self.t_tbl_cells.setRowCount(len(cells))
        for i,c in enumerate(cells):
            bbox = c.get("bbox",[])
            bbox_str = f"{bbox[0]:.0f},{bbox[1]:.0f},{bbox[2]:.0f},{bbox[3]:.0f}" if len(bbox)==4 else ""
            conf = c.get("confidence","")
            conf_str = f"{conf:.2%}" if isinstance(conf, float) else str(conf)
            for j,v in enumerate([
                c.get("row",""), c.get("col",""),
                c.get("row_span",1), c.get("col_span",1),
                c.get("text",""), conf_str, bbox_str
            ]):
                item2 = QTableWidgetItem(str(v))
                if j == 5 and isinstance(conf, float):  # 置信度着色
                    if conf >= 0.9:
                        item2.setForeground(QColor("#2f9e44"))
                    elif conf >= 0.7:
                        item2.setForeground(QColor("#e67700"))
                    else:
                        item2.setForeground(QColor("#c92a2a"))
                self.t_tbl_cells.setItem(i,j,item2)
        self.t_txt_json.setText(json.dumps(tbl,ensure_ascii=False,indent=2))

    def _t_start(self):
        if not self._t_files: return
        for b in (self.t_btn_run,self.t_btn_sel,self.t_btn_clr,self.t_btn_xl,self.t_btn_js): b.setEnabled(False)
        self.t_btn_stop.setEnabled(True)
        self.t_banner.update(0,"正在初始化 PP-StructureV3...")
        self.t_tbl_list.clear(); self.t_tbl_cells.setRowCount(0)
        self.t_browser.clear(); self.t_txt_json.clear()
        self.lbl_status.setText("正在识别表格...")
        self._t_results=[]
        self._t_worker=TableStructureWorker(self._t_files)
        self._t_worker.progress.connect(self._t_on_prog)
        self._t_worker.finished.connect(self._t_on_done)
        self._t_worker.error.connect(self._t_on_err)
        self._t_worker.start()

    def _t_stop(self):
        if self._t_worker: self._t_worker.stop()
        self.t_btn_stop.setEnabled(False); self.lbl_status.setText("正在停止...")

    def _t_on_prog(self,val,msg):
        self.t_banner.update(val,msg); self.lbl_status.setText(msg)

    def _t_on_done(self,results):
        self._t_results=results
        for b in (self.t_btn_run,self.t_btn_sel,self.t_btn_clr,self.t_btn_xl,self.t_btn_js): b.setEnabled(True)
        self.t_btn_stop.setEnabled(False); self.t_banner.done()
        self.lbl_status.setText("识别完成"); self._t_update_summary()
        ok=sum(1 for r in results if r.get("success"))
        total=sum(len(r.get("tables",[])) for r in results if r.get("success"))
        QMessageBox.information(self,"完成",f"识别完成！\n成功: {ok}/{len(results)} 个文件\n共识别 {total} 张表格")
        if self.t_file_list.count()>0:
            self.t_file_list.setCurrentRow(0); self._t_on_file_click(self.t_file_list.item(0))

    def _t_on_err(self,msg):
        for b in (self.t_btn_run,self.t_btn_sel,self.t_btn_clr): b.setEnabled(True)
        self.t_btn_stop.setEnabled(False); self.t_banner.done()
        self.lbl_status.setText("识别失败"); QMessageBox.critical(self,"错误",msg)

    def _t_update_summary(self):
        ok=sum(1 for r in self._t_results if r.get("success"))
        total=sum(len(r.get("tables",[])) for r in self._t_results if r.get("success"))
        lines=[f"文件数: {len(self._t_results)}  ( 成功 {ok} )",f"表格总数: {total}"]
        self.t_lbl_summary.setText("\n".join(lines))
