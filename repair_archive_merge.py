# -*- coding: utf-8 -*-
"""调修单批量识别表 + 返修卡批量识别表 → 返修件档案.xlsx（四键匹配，内连/左连）。"""

from __future__ import annotations

import math
import os
import tempfile
from io import BytesIO
from typing import Any, Dict, List, Literal, Optional, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ORDER_KEY_COLS = ['装备型号', '器材名称', '器件编号', '型（图）号']
CARD_KEY_COLS = ['产品代号', '返修件名称', '批次号', '图号']
UNIFY = ['_k_model', '_k_name', '_k_batch', '_k_drawing']

ARCHIVE_HEADERS = [
    '返修卡号', '产品代号', '批次号', '返修件名称', '图号',
    '返修故障件信息', '损坏原因修理结果', 'FRACAS/排故报告编号',
    '返修卡原图', '调修单号', '邮寄地址', '进厂时间', '调修单原图',
]

# 批量导出模板中「原图」列（1-based）
ORDER_IMG_COL = 10   # J
CARD_IMG_COL = 14    # N
# 档案表中嵌入列（1-based）
ARCHIVE_CARD_IMG_COL = 9   # I
ARCHIVE_ORDER_IMG_COL = 13  # M


def _norm_key(v: Any) -> str:
    if v is None:
        return ''
    if isinstance(v, float) and pd.isna(v):
        return ''
    s = str(v).strip().replace('\xa0', ' ').strip()
    return s


def _strip_df_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c).strip().replace('\xa0', ' ') if c is not None else '' for c in out.columns]
    return out


def _require_columns(df: pd.DataFrame, names: List[str], label: str) -> None:
    miss = [n for n in names if n not in df.columns]
    if miss:
        raise ValueError(f'{label} 缺少列：{", ".join(miss)}。实际表头：{list(df.columns)}')


def _excel_row_num(v: Any) -> Optional[int]:
    if v is None or (isinstance(v, float) and (math.isnan(v) or pd.isna(v))):
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _anchor_cell_0based(anchor: Any) -> Optional[Tuple[int, int]]:
    if anchor is None:
        return None
    f = getattr(anchor, '_from', None)
    if f is None:
        return None
    return int(f.row), int(f.col)


def _image_bytes_at(ws: Any, row_1based: int, col_1based: int) -> Optional[bytes]:
    r0, c0 = row_1based - 1, col_1based - 1
    for im in getattr(ws, '_images', []) or []:
        t = _anchor_cell_0based(im.anchor)
        if t == (r0, c0):
            try:
                return im._data()
            except Exception:
                continue
    return None


def _thumb_png(data: bytes, max_w: int, max_h: int) -> Optional[bytes]:
    try:
        from PIL import Image as PILImage
        pil = PILImage.open(BytesIO(data))
        if pil.mode in ('RGBA', 'P'):
            pil = pil.convert('RGB')
        try:
            resample = PILImage.Resampling.LANCZOS
        except AttributeError:
            resample = PILImage.LANCZOS  # type: ignore[attr-defined]
        pil.thumbnail((max_w, max_h), resample)
        buf = BytesIO()
        pil.save(buf, format='PNG')
        return buf.getvalue()
    except Exception:
        return None


def build_repair_archive_xlsx(
    order_xlsx_path: str,
    card_xlsx_path: str,
    out_xlsx_path: str,
    join_type: Literal['inner', 'left'] = 'inner',
) -> Dict[str, Any]:
    """
    读取两个 xlsx 首工作表，按四键合并后写出「返修件档案」表。
    join_type: inner — 仅双方键完全一致；left — 以返修卡表为主，未匹配调修单字段留空。
    """
    order_df = pd.read_excel(order_xlsx_path, sheet_name=0, dtype=object, engine='openpyxl')
    card_df = pd.read_excel(card_xlsx_path, sheet_name=0, dtype=object, engine='openpyxl')
    order_df = _strip_df_columns(order_df)
    card_df = _strip_df_columns(card_df)

    _require_columns(order_df, ORDER_KEY_COLS + ['调修单号', '邮寄地址', '进厂时间'], '调修单批量识别表')
    _require_columns(
        card_df,
        CARD_KEY_COLS + ['返修卡号', '返修故障件信息', '损坏原因修理结果', 'FRACAS/排故报告编号'],
        '返修卡批量识别表',
    )

    order_p = order_df.copy()
    card_p = card_df.copy()
    for k, u in zip(ORDER_KEY_COLS, UNIFY):
        order_p[u] = order_p[k].map(_norm_key)
    for k, u in zip(CARD_KEY_COLS, UNIFY):
        card_p[u] = card_p[k].map(_norm_key)

    order_p['_row_o'] = order_p.index + 2
    card_p['_row_c'] = card_p.index + 2

    how: Literal['inner', 'left'] = 'inner' if join_type == 'inner' else 'left'
    merged = pd.merge(card_p, order_p, on=UNIFY, how=how, suffixes=('_c_only', '_o_only'))

    dup_keys = int((merged.groupby(UNIFY).size() > 1).sum()) if len(merged) else 0

    wb_o = load_workbook(order_xlsx_path, data_only=False)
    wb_c = load_workbook(card_xlsx_path, data_only=False)
    ws_o = wb_o.active
    ws_c = wb_c.active

    wb = Workbook()
    ws = wb.active
    ws.title = '返修件档案'

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=10)
    thin = Side(style='thin', color='B4C6E7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align = Alignment(vertical='center', wrap_text=True)

    for col, title in enumerate(ARCHIVE_HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border

    widths = [12, 10, 12, 14, 14, 30, 30, 18, 16, 22, 28, 12, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    card_fields = [
        '返修卡号', '产品代号', '批次号', '返修件名称', '图号',
        '返修故障件信息', '损坏原因修理结果', 'FRACAS/排故报告编号',
    ]
    order_fields = ['调修单号', '邮寄地址', '进厂时间']

    def _cell_str(v: Any) -> str:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return ''
        return str(v)

    temp_pngs: List[str] = []
    try:
        row_idx = 2
        for _, row in merged.iterrows():
            for col, fn in enumerate(card_fields, start=1):
                cell = ws.cell(row=row_idx, column=col, value=_cell_str(row.get(fn, '')))
                cell.alignment = align
                cell.border = border

            for j, fn in enumerate(order_fields):
                col = 10 + j
                cell = ws.cell(row=row_idx, column=col, value=_cell_str(row.get(fn, '')))
                cell.alignment = align
                cell.border = border

            ro = _excel_row_num(row.get('_row_o'))
            rc = _excel_row_num(row.get('_row_c'))

            ws.row_dimensions[row_idx].height = 28

            if rc is not None:
                raw = _image_bytes_at(ws_c, rc, CARD_IMG_COL)
                if raw:
                    png = _thumb_png(raw, 480, 360)
                    if png:
                        tp = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
                        tp.write(png)
                        tp.close()
                        temp_pngs.append(tp.name)
                        xl_img = XLImage(tp.name)
                        ws.add_image(xl_img, f'{get_column_letter(ARCHIVE_CARD_IMG_COL)}{row_idx}')
                        ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height or 28, 200)

            if ro is not None:
                raw = _image_bytes_at(ws_o, ro, ORDER_IMG_COL)
                if raw:
                    png = _thumb_png(raw, 520, 380)
                    if png:
                        tp = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
                        tp.write(png)
                        tp.close()
                        temp_pngs.append(tp.name)
                        xl_img = XLImage(tp.name)
                        ws.add_image(xl_img, f'{get_column_letter(ARCHIVE_ORDER_IMG_COL)}{row_idx}')
                        ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height or 28, 220)

            row_idx += 1

        wb.save(out_xlsx_path)
    finally:
        wb_o.close()
        wb_c.close()
        for p in temp_pngs:
            try:
                os.unlink(p)
            except OSError:
                pass

    return {
        'rows_out': len(merged),
        'rows_order': len(order_df),
        'rows_card': len(card_df),
        'join_type': join_type,
        'duplicate_key_groups': dup_keys,
    }
